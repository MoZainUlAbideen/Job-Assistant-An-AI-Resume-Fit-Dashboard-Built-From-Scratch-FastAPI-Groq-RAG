"""
File -> text, plus the layout signals an ATS parser actually trips over.

Two jobs here:
  1. Get the text out (pdfplumber for PDF, python-docx for DOCX).
  2. Record *how* the text was laid out, because that's what makes resumes
     fail in real ATS parsers even when they look fine to a human:
       - text inside tables / text boxes / headers & footers (often skipped)
       - two-column PDF layouts (reading order gets interleaved)
       - "(cid:123)" glyphs (custom fonts with no unicode map -> gibberish)
       - scanned/image-only PDFs (no text layer at all)
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

import docx
import pdfplumber

MAX_BYTES = 5 * 1024 * 1024


@dataclass
class Extracted:
    text: str
    file_type: str                       # "pdf" | "docx" | "txt"
    page_count: int = 1
    has_tables: bool = False
    table_text_chars: int = 0
    header_footer_text: str = ""
    has_text_boxes: bool = False
    cid_glyphs: int = 0
    multi_column_pages: int = 0
    docx_paragraphs: list[str] = field(default_factory=list)  # body paragraphs in order (docx only)


def extract(file_bytes: bytes, filename: str) -> Extracted:
    if len(file_bytes) > MAX_BYTES:
        raise ValueError("File is larger than 5 MB.")
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _from_pdf(file_bytes)
    if name.endswith(".docx"):
        return _from_docx(file_bytes)
    if name.endswith(".txt") or name.endswith(".md"):
        return Extracted(text=file_bytes.decode("utf-8", errors="replace"), file_type="txt")
    raise ValueError("Unsupported file type — upload a .pdf, .docx or .txt resume.")


# ---------------------------------------------------------------- PDF
def _from_pdf(file_bytes: bytes) -> Extracted:
    chunks: list[str] = []
    multi_col = 0
    try:
        pdf = pdfplumber.open(BytesIO(file_bytes))
    except Exception as e:  # pdfminer raises a zoo of exception types on corrupt files
        raise ValueError(f"Could not read this PDF ({type(e).__name__}).") from e
    with pdf:
        pages = len(pdf.pages)
        for page in pdf.pages:
            t = page.extract_text(x_tolerance=1.5, y_tolerance=3) or ""
            chunks.append(t)
            if _looks_multi_column(page):
                multi_col += 1
    text = "\n".join(chunks)
    return Extracted(
        text=_clean(text), file_type="pdf", page_count=pages,
        cid_glyphs=len(re.findall(r"\(cid:\d+\)", text)), multi_column_pages=multi_col,
        has_tables=False,
    )


def _looks_multi_column(page) -> bool:
    """Heuristic: a two-column layout has many lines where a wide horizontal gap
    separates two *substantial* runs of words. Single-column resumes also have
    gaps (right-aligned dates/locations), so we only count a line when the
    right-hand run is long (>= 6 words) -- "January 2022 – December 2023" is 5."""
    try:
        words = page.extract_words(x_tolerance=1.5, y_tolerance=3)
    except Exception:
        return False
    if len(words) < 40:
        return False
    width = float(page.width)
    lines: dict[int, list] = {}
    for w in words:
        lines.setdefault(round(float(w["top"]) / 3), []).append(w)
    split_lines = 0
    for ws in lines.values():
        ws.sort(key=lambda w: w["x0"])
        for i in range(1, len(ws)):
            gap = float(ws[i]["x0"]) - float(ws[i - 1]["x1"])
            if gap > 0.12 * width and len(ws) - i >= 6:
                split_lines += 1
                break
    return split_lines >= max(6, 0.25 * len(lines))


# ---------------------------------------------------------------- DOCX
def _from_docx(file_bytes: bytes) -> Extracted:
    try:
        d = docx.Document(BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Could not read this DOCX ({type(e).__name__}).") from e

    paras = [p.text for p in d.paragraphs]
    # Word bullets are paragraph formatting (a "List" style or numbering), not a "•" in the text.
    # Mark them so the resume parser can tell bullets from role/project header lines.
    marked = []
    for p in d.paragraphs:
        is_list = p.style is not None and p.style.name.lower().startswith("list") \
            or p._p.pPr is not None and p._p.pPr.numPr is not None
        marked.append(("• " + p.text) if (is_list and p.text.strip() and not p.text.lstrip().startswith("•")) else p.text)
    table_lines: list[str] = []
    for t in d.tables:
        for row in t.rows:
            seen = set()
            cells = []
            for c in row.cells:            # merged cells repeat; de-dup by identity
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                if c.text.strip():
                    cells.append(c.text.strip())
            if cells:
                table_lines.append("\n".join(cells))

    hf: list[str] = []
    for s in d.sections:
        for part in (s.header, s.footer):
            try:
                hf.extend(p.text for p in part.paragraphs if p.text.strip())
            except Exception:
                pass

    body_xml = d.element.body.xml
    has_text_boxes = "txbxContent" in body_xml

    text = "\n".join([p for p in marked if p.strip()] + table_lines)
    return Extracted(
        text=_clean(text), file_type="docx",
        has_tables=bool(table_lines), table_text_chars=sum(len(x) for x in table_lines),
        header_footer_text="\n".join(hf), has_text_boxes=has_text_boxes,
        docx_paragraphs=paras,
    )


def _clean(text: str) -> str:
    # bullet glyphs from fonts with no unicode map come out as "(cid:127)"; they're counted
    # for the format check before this normalises the line-leading ones to a real bullet
    text = re.sub(r"(?m)^[ \t]*\(cid:\d+\)[ \t]*", "• ", text)
    text = text.replace(" ", " ").replace("", "•").replace("•", "•")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
