"""Shared fixtures: build DOCX/PDF resumes on the fly from the synthetic text fixture."""
from io import BytesIO
from pathlib import Path

import docx
import pytest

FIX = Path(__file__).parent / "fixtures"
RESUME_TXT = (FIX / "resume_ayesha.txt").read_text(encoding="utf-8")
JD_TXT = (FIX / "jd_mle.txt").read_text(encoding="utf-8")


def make_docx(text: str = RESUME_TXT, contact_in_header: bool = False, table: bool = False) -> bytes:
    d = docx.Document()
    lines = text.splitlines()
    if contact_in_header:
        d.sections[0].header.paragraphs[0].text = lines[2]
        lines = lines[:2] + lines[3:]
    for line in lines:
        if line.startswith("• "):
            d.add_paragraph(line[2:], style="List Bullet")
        else:
            d.add_paragraph(line)
    if table:
        t = d.add_table(rows=6, cols=2)
        for r in t.rows:
            r.cells[0].text = "Extra skills listed in a table cell to make this table large " * 3
            r.cells[1].text = "More content in a second column " * 3
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(text: str = RESUME_TXT, two_column: bool = False) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    lines = [l for l in text.splitlines()]
    if two_column:
        filler = "Long sentence of words in this column for layout testing purposes here"
        y = h - 50
        for i in range(45):
            c.drawString(40, y, (lines[i % len(lines)] or "x")[:45])
            c.drawString(w / 2 + 20, y, filler[:60])
            y -= 15
    else:
        y = h - 40
        for l in lines:
            c.drawString(40, y, l[:110])
            y -= 13
            if y < 40:
                c.showPage()
                y = h - 40
    c.save()
    return buf.getvalue()


@pytest.fixture
def fake_llm(monkeypatch):
    """Replace the Groq call with canned proposals (one honest, two that must be blocked)."""
    from app import llm

    calls = []

    def chat_json(system, user, max_tokens=3000):
        import json
        calls.append(json.loads(user))
        bullets = {b["text"][:40]: b["id"] for b in calls[-1]["bullets"]}
        pick = lambda prefix: next(v for k, v in bullets.items() if k.startswith(prefix))
        return {"edits": [
            {"id": pick("Built a document classification"),
             "text": "Built a deep learning document classification service in PyTorch and FastAPI serving 2M requests per month.",
             "keywords": ["Deep Learning"]},
            {"id": pick("Containerized the API"),
             "text": "Containerized the API with Docker and deployed it on Kubernetes and a Linux VPS.",
             "keywords": ["Kubernetes"]},
            {"id": pick("Implemented a retrieval"),
             "text": "Implemented a retrieval pipeline over 5,000 resumes using FAISS vector search and Groq-hosted Llama 3, implying agent-like LLMs.",
             "keywords": ["LLMs"]},
        ]}

    monkeypatch.setenv("GROQ_API_KEY", "test-key")
    monkeypatch.setattr(llm, "chat_json", chat_json)
    return calls
